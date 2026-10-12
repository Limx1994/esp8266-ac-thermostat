from pathlib import Path
import csv, hashlib, json, zipfile, warnings
from pypdf import PdfReader
from openpyxl import load_workbook

root = Path('D:/空调温度')
out = root / '审查报告/20261011'
baseline = json.loads((out / 'recovery_baseline.json').read_text(encoding='utf-8-sig'))
results = []
for item in baseline:
    path = root / item['Path']
    ext = path.suffix.lower()
    if item['Path'].startswith('Sof\\build\\') or ext not in ('.pdf', '.xlsx', '.csv', '.zip'):
        continue
    record = {'path': str(path), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
              'baseline_sha256': item['SHA256'], 'attempts': 1,
              'limit': '文本/单元格/结构读取，不证明图形、电气连接或当前固件实机行为'}
    try:
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter('always')
            if ext == '.pdf':
                reader = PdfReader(path)
                texts = [page.extract_text() or '' for page in reader.pages]
                target = out / ('attachment_' + path.stem + '.txt')
                target.write_text('\n'.join(f'PAGE {i+1}\n{text}' for i, text in enumerate(texts)), encoding='utf-8')
                record.update(pages=len(texts), text_chars=sum(map(len, texts)),
                              empty_text_pages=[i+1 for i, text in enumerate(texts) if not text.strip()],
                              extraction_path=str(target), excerpt=' '.join(texts[0].split())[:600])
            elif ext == '.xlsx':
                book = load_workbook(path, read_only=True, data_only=False)
                sheets = {sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)] for sheet in book}
                book.close()
                target = out / ('attachment_' + path.stem + '_cells.json')
                target.write_text(json.dumps(sheets, ensure_ascii=False, default=str, indent=2), encoding='utf-8')
                record.update(sheets={name: len(rows) for name, rows in sheets.items()}, extraction_path=str(target))
            elif ext == '.csv':
                with path.open(encoding='utf-8-sig', newline='') as stream:
                    rows = list(csv.reader(stream))
                record.update(rows=len(rows), header=rows[0] if rows else [],
                              column_counts=sorted({len(row) for row in rows}),
                              first_data=rows[1] if len(rows)>1 else [], last_data=rows[-1] if rows else [])
            else:
                with zipfile.ZipFile(path) as archive:
                    record['members'] = [{'name': member.filename, 'bytes': member.file_size,
                                          'crc': member.CRC} for member in archive.infolist()]
            record['warnings'] = [str(w.message) for w in caught]
        record['status'] = 'current_bytes_read'
    except Exception as exc:
        record.update(status='limited', error=f'{type(exc).__name__}: {exc}')
    record['hash_matches_baseline'] = record['sha256'].upper() == record['baseline_sha256'].upper()
    results.append(record)
(out / 'completeness_attachments_current.json').write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps([{'path': r['path'], 'status': r['status'], 'pages': r.get('pages'), 'rows': r.get('rows'), 'error': r.get('error')} for r in results], ensure_ascii=False))
